# -*- coding: utf-8 -*-
"""Excel 周报导出：严格对齐用户模板《2026年东城区新闻发稿统计.xlsx》。

模板结构（Sheet1）：
  第1行  A1:J1 合并大标题（宋体24加粗，行高30）
  第2行  表头：序号 | 日期 | 媒体名称 | (标题) | 版面 | 位置 |
        涉及单位1 | 涉及单位2 | 涉及单位3 | 关键词1（宋体12加粗，行高24）
  数据行：行高24、全边框、序号 =ROW()-2 公式
  列宽  ：B≈9.8 C≈59 D≈142.3 E≈7.8 F≈8.2 G≈27.8 H≈25 I≈28.8 J≈14.2
"""
import os
import re
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from . import config, db

HEADERS = ["序号", "日期", "媒体名称", None, "版面", "位置",
           "涉及单位1", "涉及单位2", "涉及单位3", "关键词1"]
COL_WIDTHS = {"A": 6.5, "B": 9.8, "C": 59, "D": 142.3, "E": 7.8,
              "F": 8.2, "G": 27.8, "H": 25, "I": 28.8, "J": 14.2}
THIN = Border(*[Side(style="thin")] * 4)

MEDIA_FALLBACK = {
    "cctv": "央视网", "chinadaily": "中国日报", "thepaper": "澎湃新闻",
    "people_bj": "人民网", "xinhua_bj": "新华网", "chinanews": "中新网",
    "qianlong": "千龙网", "bjnews": "新京报", "cri": "国际在线",
    "bjdc_gov": "东城区人民政府网",
}


def _parse_date(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def fetch_weekly_rows(days=7):
    """最近 days 天（含今日）的命中新闻，按发布日期升序。"""
    today = datetime.now().date()
    d_from = (today - timedelta(days=days - 1)).strftime("%Y-%m-%d")
    d_to = today.strftime("%Y-%m-%d")
    rows = db.query_by_date_range(d_from, d_to)
    # 无发布日期的条目：按入库时间归入窗口则保留（用抓取日期兜底）
    no_date = db.query_no_date()
    fetched_in_window = []
    for r in no_date:
        fa = (r["fetched_at"] or "")[:10]
        if d_from <= fa <= d_to:
            fetched_in_window.append(dict(r, pub_date=fa))
    rows = [dict(r) for r in rows] + fetched_in_window
    rows.sort(key=lambda r: (r["pub_date"] or "9999", r["media"] or "",
                             r["title"] or ""))
    return rows, d_from, d_to


def build_sheet(rows, d_from, d_to, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # 第1行：合并标题
    ws.merge_cells("A1:J1")
    title = title or (f"东城区新闻发稿统计（{d_from}至{d_to}）")
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name="宋体", size=24, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # 第2行：表头
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=2, column=col, value=h)
        c.font = Font(name="宋体", size=12, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN
    ws.row_dimensions[2].height = 24

    # 数据行
    for i, r in enumerate(rows):
        excel_row = i + 3
        ws.row_dimensions[excel_row].height = 24
        d = _parse_date(r.get("pub_date"))
        units = (r.get("units") or "").split("、")
        units += [""] * (3 - len(units))
        # 版面：电子报源的摘要形如“【版面】第3版北京新闻｜正文……”
        banmian = ""
        m = re.match(r"【版面】([^｜|]+)", r.get("summary") or "")
        if m:
            banmian = m.group(1).strip()
        values = [
            f"=ROW()-2",                                   # 序号（与模板一致）
            d,                                             # 日期
            r.get("media") or MEDIA_FALLBACK.get(
                r.get("source_key"), r.get("source_key") or "未知来源"),
            r.get("title") or "",                          # 标题
            banmian or None,                               # 版面（电子报填写）
            None,                                          # 位置
            units[0] or None, units[1] or None, units[2] or None,
            r.get("category") or "综合",
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=excel_row, column=col, value=v)
            c.font = Font(name="宋体", size=12)
            c.border = THIN
            if col == 2 and d is not None:
                c.number_format = "yyyy/m/d"
                c.alignment = Alignment(horizontal="center",
                                        vertical="center")
            elif col == 1:
                c.alignment = Alignment(horizontal="center",
                                        vertical="center")
            else:
                c.alignment = Alignment(vertical="center", wrap_text=False)

    # 列宽（与模板一致）
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w
    return wb


def export_weekly(days=7, out_path=None, quiet=False):
    """生成最近一周的统计 Excel。返回文件路径。"""
    db.init_db()
    rows, d_from, d_to = fetch_weekly_rows(days)
    if not rows:
        print(f"最近 {days} 天内暂无命中新闻，"
              f"可先运行 `start` 或 `crawl` 抓取数据。")
        return None
    wb = build_sheet(rows, d_from, d_to)
    if not out_path:
        out_path = os.path.join(
            config.EXPORT_DIR,
            f"东城区新闻发稿统计_{d_from.replace('-', '')}"
            f"-{d_to.replace('-', '')}.xlsx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    if not quiet:
        print(f"已导出：{out_path}")
        print(f"统计区间：{d_from} ~ {d_to}，共 {len(rows)} 条")
        by_media = {}
        for r in rows:
            k = r.get("media") or r.get("source_key")
            by_media[k] = by_media.get(k, 0) + 1
        for k, v in sorted(by_media.items(), key=lambda x: -x[1]):
            print(f"  {k}: {v} 条")
    return out_path
